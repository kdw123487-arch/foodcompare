#!/usr/bin/env python3
"""data/products.xlsx  ->  data/products.json

사이트(index.html)는 products.json만 읽습니다. 제품·항목·박스 구성은 모두 엑셀에서 고치고,
이 스크립트로 JSON을 다시 만듭니다.   사용법:  python3 tools/build_data.py
오류가 있으면 시트 이름과 줄 번호를 알려주고 JSON은 만들지 않습니다.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "products.xlsx"
DST = ROOT / "data" / "products.json"

errors = []


def err(sheet, row, msg):
    errors.append(f"[{sheet}] {row}번째 줄: {msg}" if row else f"[{sheet}] {msg}")


def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def to_num(v, sheet, row, col):
    """숫자 칸: 빈 칸 → None, 숫자 → 숫자, 그 외 → 오류"""
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, bool):
        err(sheet, row, f"'{col}' 칸에 숫자가 아닌 값이 있습니다: {v!r}")
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        f = float(str(v).replace(",", ""))
        return int(f) if f.is_integer() else f
    except ValueError:
        err(sheet, row, f"'{col}' 칸에 숫자가 아닌 값이 있습니다: {v!r}")
        return None


def table(wb, name, columns, required=(), extra_ok=False):
    """시트를 [(줄번호, {열이름: 값})]로 읽는다. columns에 없는 머리글은 오류(extra_ok면 허용)."""
    if name not in wb.sheetnames:
        err(name, None, "시트가 없습니다")
        return [], []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [clean(h) for h in rows[0]]
    for h in header:
        if h and h not in columns and not extra_ok:
            err(name, 1, f"알 수 없는 열 이름 '{h}' (쓸 수 있는 열: {', '.join(columns)})")
    for h in required:
        if h not in header:
            err(name, 1, f"필수 열 '{h}'이(가) 없습니다")
    out = []
    for i, row in enumerate(rows[1:], start=2):
        d = {h: clean(v) for h, v in zip(header, row) if h}
        if any(v is not None for v in d.values()):
            out.append((i, d))
    return out, header


def split_keys(text):
    return [k.strip() for k in re.split(r"[,\s]+", text or "") if k.strip()]


def parse_trio(text, sheet, row):
    cols = []
    for part in (text or "").split("|"):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"([A-Za-z0-9_]+)\s*(?:\[([^\]]*)\])?", part)
        if not m:
            err(sheet, row, f"trio 항목 문법 오류: '{part}' (예: carb[sugar] | protein | fat[satFat,transFat])")
            continue
        cols.append({"field": m.group(1), "subs": split_keys(m.group(2))})
    return cols


def main():
    if not SRC.exists():
        sys.exit(f"엑셀 파일이 없습니다: {SRC}")
    wb = load_workbook(SRC, data_only=True, read_only=True)

    # ---------- 사이트 ----------
    site = {}
    for i, d in table(wb, "사이트", ["키", "값"], required=["키", "값"])[0]:
        if d.get("키"):
            site[d["키"]] = d.get("값")

    # ---------- 항목 ----------
    fields, label_to_key = {}, {}
    for i, d in table(wb, "항목", ["키", "이름", "단위", "소수 최소", "소수 최대"], required=["키", "이름"])[0]:
        k = d.get("키")
        if not k:
            err("항목", i, "키가 비어 있습니다")
            continue
        if k in fields:
            err("항목", i, f"키 '{k}'가 중복입니다")
            continue
        lo = to_num(d.get("소수 최소"), "항목", i, "소수 최소") or 0
        hi = to_num(d.get("소수 최대"), "항목", i, "소수 최대")
        hi = lo if hi is None else hi
        fields[k] = {"label": d.get("이름") or k, "unit": d.get("단위") or "", "dec": [int(lo), int(max(lo, hi))]}
        label_to_key[fields[k]["label"]] = k

    # ---------- 카테고리 ----------
    cat_cols = ["id", "이름", "라벨", "설명", "빈 화면 안내", "기준 항목", "요약 항목", "머리글 항목"]
    categories = []
    for i, d in table(wb, "카테고리", cat_cols, required=["id", "이름"])[0]:
        categories.append({
            "id": d["id"], "name": d["이름"], "label": d.get("라벨") or "", "desc": d.get("설명") or "",
            "empty": d.get("빈 화면 안내") or "", "ref": d.get("기준 항목") or None,
            "summary": split_keys(d.get("요약 항목")), "header": split_keys(d.get("머리글 항목")), "boxes": [],
        })
    cat_by_id = {c["id"]: c for c in categories}

    # ---------- 박스 ----------
    boxes = []
    for i, d in table(wb, "박스", ["카테고리", "순서", "종류", "제목", "부제", "항목"], required=["카테고리", "종류"])[0]:
        c = cat_by_id.get(d.get("카테고리"))
        if not c:
            err("박스", i, f"카테고리 '{d.get('카테고리')}'가 카테고리 시트에 없습니다")
            continue
        order = to_num(d.get("순서"), "박스", i, "순서") or 0
        kind, items = d.get("종류"), d.get("항목") or ""
        if kind == "trio":
            box = {"type": "trio", "columns": parse_trio(items, "박스", i)}
        elif kind == "sum":
            box = {"type": "sum", "title": d.get("제목") or "", "parts": split_keys(items)}
        elif kind == "list":
            box = {"type": "list", "title": d.get("제목") or "", "note": d.get("부제") or "", "source": items.strip()}
        else:
            err("박스", i, f"종류는 trio / sum / list 중 하나여야 합니다: '{kind}'")
            continue
        boxes.append((c["id"], order, i, box))
    for cid, _, _, box in sorted(boxes, key=lambda t: (t[0], t[1], t[2])):
        cat_by_id[cid]["boxes"].append(box)

    # ---------- 제품 ----------
    p_cols = ["제품id", "카테고리", "브랜드", "제품명", "총중량(g)", "1회섭취량(g)", "총제공횟수(회)", "가격(원)", "이미지", "출처", "확인일", "기준맛"]
    products, prod_by_id = [], {}
    for i, d in table(wb, "제품", p_cols, required=["제품id", "카테고리", "제품명"])[0]:
        pid = d.get("제품id")
        if pid in prod_by_id:
            err("제품", i, f"제품id '{pid}'가 중복입니다")
            continue
        if d.get("카테고리") not in cat_by_id:
            err("제품", i, f"카테고리 '{d.get('카테고리')}'가 카테고리 시트에 없습니다")
        checked = d.get("확인일")
        if hasattr(checked, "strftime"):
            checked = checked.strftime("%Y.%m.%d")
        # 이미지 칸: 사진 경로 여러 장은 | 로 구분 (슬라이드로 넘겨 볼 수 있음)
        images = [x.strip() for x in re.split(r"[|\n]", d.get("이미지") or "") if x.strip()]
        for im in images:
            if not (ROOT / im).exists():
                err("제품", i, f"사진 파일을 찾을 수 없습니다: {im}")
        p = {
            "id": pid, "category": d.get("카테고리"), "brand": d.get("브랜드") or "", "name": d["제품명"],
            "weight": to_num(d.get("총중량(g)"), "제품", i, "총중량(g)"),
            "serving": to_num(d.get("1회섭취량(g)"), "제품", i, "1회섭취량(g)"),
            "servings": to_num(d.get("총제공횟수(회)"), "제품", i, "총제공횟수(회)"),
            "price": to_num(d.get("가격(원)"), "제품", i, "가격(원)"),
            "images": images, "source": d.get("출처") or "", "checked": str(checked) if checked else "",
            "baseFlavor": d.get("기준맛"), "lists": {}, "flavors": [],
        }
        products.append(p)
        prod_by_id[pid] = p

    # ---------- 맛 ----------
    f_fixed = {"제품id", "맛id", "맛", "1회섭취량(g)", "가격(원)", "링크"}
    rows, header = table(wb, "맛", f_fixed, required=["제품id", "맛id", "맛"], extra_ok=True)
    col_key = {}
    for h in header:
        if h and h not in f_fixed:
            k = label_to_key.get(h) or (h if h in fields else None)
            if k is None:
                err("맛", 1, f"열 이름 '{h}'이(가) 항목 시트에 없습니다. 항목 시트에 먼저 추가하세요 (쓸 수 있는 이름: {', '.join(label_to_key)})")
            else:
                col_key[h] = k
    seen = set()
    for i, d in rows:
        p = prod_by_id.get(d.get("제품id"))
        if not p:
            err("맛", i, f"제품id '{d.get('제품id')}'가 제품 시트에 없습니다")
            continue
        fid = d.get("맛id")
        if fid in seen:
            err("맛", i, f"맛id '{fid}'가 중복입니다")
        seen.add(fid)
        values = {}
        for h, k in col_key.items():
            v = to_num(d.get(h), "맛", i, h)
            if v is not None:
                values[k] = v
        fl = {"id": fid, "name": d.get("맛"), "values": values}
        sv = to_num(d.get("1회섭취량(g)"), "맛", i, "1회섭취량(g)")
        if sv:
            fl["serving"] = sv   # 맛마다 1회 섭취량이 다를 때만 (비우면 제품 값 사용)
        pr = to_num(d.get("가격(원)"), "맛", i, "가격(원)")
        if pr:
            fl["price"] = pr   # 맛마다 가격이 다를 때만 (비우면 제품 가격 사용)
        lk = (d.get("링크") or "").strip()
        if lk:
            if not lk.startswith(("http://", "https://")):
                err("맛", i, "링크는 http:// 또는 https:// 로 시작해야 합니다")
            else:
                fl["link"] = lk   # 맛마다 판매 페이지가 다를 때만 (2개 이상이면 사이트에서 목록으로 펼쳐짐)
        p["flavors"].append(fl)

    # ---------- 목록 ----------
    # 맛id가 있으면 그 맛 전용 목록, 비어 있으면 제품 전체 기본 목록. 부제 칸은 박스 머리 문구(목록당 하나).
    for i, d in table(wb, "목록", ["제품id", "맛id", "목록", "이름", "설명", "값", "단위", "부제"], required=["제품id", "목록", "이름"])[0]:
        p = prod_by_id.get(d.get("제품id"))
        if not p:
            err("목록", i, f"제품id '{d.get('제품id')}'가 제품 시트에 없습니다")
            continue
        owner = p
        if d.get("맛id"):
            owner = next((f for f in p["flavors"] if f["id"] == d["맛id"]), None)
            if owner is None:
                err("목록", i, f"맛id '{d['맛id']}'가 제품 '{p['id']}'의 맛 시트에 없습니다")
                continue
        owner.setdefault("lists", {}).setdefault(d["목록"], []).append(
            {"name": d["이름"], "desc": d.get("설명") or "", "value": to_num(d.get("값"), "목록", i, "값"), "unit": d.get("단위") or ""})
        if d.get("부제"):
            notes = owner.setdefault("listNotes", {})
            if notes.get(d["목록"], d["부제"]) != d["부제"]:
                err("목록", i, f"같은 목록 '{d['목록']}'에 서로 다른 부제가 있습니다")
            notes[d["목록"]] = d["부제"]

    # ---------- 참조 검사 ----------
    def need(key, where):
        if key not in fields:
            err(where[0], None, f"{where[1]}에서 쓴 항목 키 '{key}'가 항목 시트에 없습니다")

    for c in categories:
        for k in c["summary"] + c["header"]:
            need(k, ("카테고리", f"'{c['id']}'의 요약/머리글 항목"))
        if c["ref"]:
            need(c["ref"], ("카테고리", f"'{c['id']}'의 기준 항목"))
        for b in c["boxes"]:
            if b["type"] == "trio":
                for col in b["columns"]:
                    for k in [col["field"]] + col["subs"]:
                        need(k, ("박스", f"'{c['id']}' trio"))
            elif b["type"] == "sum":
                for k in b["parts"]:
                    need(k, ("박스", f"'{c['id']}' sum '{b['title']}'"))
    for p in products:
        if p["baseFlavor"] and p["baseFlavor"] not in [f["id"] for f in p["flavors"]]:
            err("제품", None, f"'{p['id']}'의 기준맛 '{p['baseFlavor']}'이(가) 맛 시트에 없습니다")
        c = cat_by_id.get(p["category"])
        if c:
            for b in c["boxes"]:
                if b["type"] == "list" and b["source"] not in p["lists"]:
                    pass  # 목록이 없는 제품은 해당 박스를 그리지 않음 (오류 아님)

    if errors:
        print("엑셀에 고칠 곳이 있습니다. JSON은 만들지 않았습니다:\n")
        print("\n".join(" - " + e for e in errors))
        sys.exit(1)

    out = {
        "_generated": "이 파일은 data/products.xlsx에서 tools/build_data.py로 만들어집니다. 직접 고치지 말고 엑셀을 고치세요.",
        "site": site, "fields": fields, "categories": categories, "products": products,
    }
    DST.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    nf = sum(len(p["flavors"]) for p in products)
    print(f"완료: 항목 {len(fields)}개 · 카테고리 {len(categories)}개 · 제품 {len(products)}개 · 맛 {nf}개 → {DST.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
